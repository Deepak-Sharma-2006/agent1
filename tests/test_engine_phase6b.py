"""
Test Suite — Presentation Engine Phase 6B: Championship Visual Primitives & Vector Icon Library
Verifies:
1. Native vector icons render across all standard types (shield, database, cloud, cpu, brain, chart, check, cross).
2. All 6 championship visual primitives compile cleanly into python-pptx presentations.
3. Shape collections are populated with native vector geometry (zero raster dependencies).
"""

import os
import unittest
from pptx import Presentation
from pptx.util import Inches

from scripts.engine.theme_registry import ThemeRegistry
from scripts.engine.flex_grid_solver import Rect
from scripts.engine.primitives import (
    draw_vector_icon,
    render_stat_hero_card,
    render_bento_card,
    render_swimlane_architecture,
    render_radial_ecosystem,
    render_tension_split_card,
    render_milestone_roadmap
)


class TestEnginePhase6B(unittest.TestCase):
    """Unit tests for Phase 6B vector icons and visual diagram primitives."""

    def setUp(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank_layout = self.prs.slide_layouts[6]
        self.tokens = ThemeRegistry.get_extended_theme("cyber_dark_terminal")
        self.theme = self.tokens.theme

    def test_vector_icons_generation(self):
        """Verifies vector icons are successfully added as native shapes to a slide."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        initial_count = len(slide.shapes)

        icons = ["shield", "database", "cloud", "cpu", "lock", "chart", "network", "brain", "check", "cross", "user"]
        for idx, icon_name in enumerate(icons):
            draw_vector_icon(slide, icon_name, left_in=1.0 + idx * 0.8, top_in=1.0, size_in=0.5)

        self.assertGreater(len(slide.shapes), initial_count + len(icons))

    def test_stat_hero_card_rendering(self):
        """Verifies stat hero card renders with number, trend pill, and label."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        rect = Rect(1.0, 1.0, 3.5, 2.4)
        render_stat_hero_card(
            slide=slide,
            rect=rect,
            stat_number="99.4%",
            stat_label="Tor De-anonymization Precision",
            delta_pill="+4.8% vs Baseline",
            caption="Verified across 10,000 live onion circuits",
            icon_type="chart",
            theme=self.theme,
            tokens=self.tokens
        )
        self.assertGreater(len(slide.shapes), 3)

    def test_bento_card_rendering(self):
        """Verifies bento card renders with title, badge pill, and formatted bullets."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        rect = Rect(1.0, 1.0, 4.0, 3.0)
        render_bento_card(
            slide=slide,
            rect=rect,
            title="Forensic Correlation Engine",
            badge="Tier 2 Core",
            bullets=[
                "Graph Intelligence: Graph Neural Networks on Neo4j cluster",
                "Temporal Resolution: Real-time packet timing correlation (<42ms)",
                "Cryptographic Proof: SHA-256 chain-of-custody logging"
            ],
            icon_type="cpu",
            theme=self.theme,
            tokens=self.tokens
        )
        self.assertGreater(len(slide.shapes), 3)

    def test_swimlane_architecture_rendering(self):
        """Verifies swimlane architecture renders multi-tier horizontal system lanes."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        rect = Rect(1.0, 1.0, 11.3, 5.0)
        lanes = [
            {"name": "Data Ingestion", "icon": "cloud", "nodes": ["Tor Relays", "Darknet Crawler", "Mempool Feeder"]},
            {"name": "Intelligence Core", "icon": "brain", "nodes": ["OnionScan Engine", "GNN Entity Resolution", "Traffic Analyzer"]},
            {"name": "Forensic Vault", "icon": "database", "nodes": ["Neo4j Cluster", "ChromaDB Vectors", "PostgreSQL"]}
        ]
        render_swimlane_architecture(slide=slide, rect=rect, lanes=lanes, theme=self.theme, tokens=self.tokens)
        self.assertGreater(len(slide.shapes), 10)

    def test_radial_ecosystem_rendering(self):
        """Verifies radial hub-and-spoke diagram renders center core and satellites."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        rect = Rect(1.0, 1.0, 11.0, 5.2)
        satellites = [
            {"name": "Law Enforcement", "desc": "Inter-agency API"},
            {"name": "Financial Intelligence", "desc": "Hawala Tracking"},
            {"name": "Cyber Defense", "desc": "Threat Hunting"},
            {"name": "Courtroom Export", "desc": "Section 63 BSA"}
        ]
        render_radial_ecosystem(
            slide=slide,
            rect=rect,
            center_title="BHEDAK CORE",
            satellites=satellites,
            center_icon="brain",
            theme=self.theme,
            tokens=self.tokens
        )
        self.assertGreater(len(slide.shapes), 8)

    def test_tension_split_card_rendering(self):
        """Verifies dual problem vs solution tension cards render with contrasting styling."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        rect = Rect(1.0, 1.0, 11.0, 4.5)
        render_tension_split_card(
            slide=slide,
            rect=rect,
            problem_title="Current Dark Web Investigation Reality",
            problem_bullets=[
                "Manual multi-hop forensic inspection takes days",
                "Tor multi-layer onion routing conceals origin server IP",
                "Evidence rejected in court due to uncertified chain-of-custody"
            ],
            solution_title="BHEDAK Autonomous Forensic Platform",
            solution_bullets=[
                "Automated sub-50ms correlation across relays",
                "Censys & Shodan passive banner fingerprinting exposes IP",
                "Cryptographic audit trail compliant with Section 63 BSA"
            ],
            theme=self.theme,
            tokens=self.tokens
        )
        self.assertGreater(len(slide.shapes), 6)

    def test_milestone_roadmap_rendering(self):
        """Verifies milestone roadmap renders horizontal progress track with status pins."""
        slide = self.prs.slides.add_slide(self.blank_layout)
        rect = Rect(1.0, 1.0, 11.0, 4.0)
        phases = [
            {"title": "Phase 1: Ingestion", "date": "Q1 2026", "status": "COMPLETED", "deliverable": "Crawler live"},
            {"title": "Phase 2: Correlation", "date": "Q2 2026", "status": "IN PROGRESS", "deliverable": "GNN model"},
            {"title": "Phase 3: Deployment", "date": "Q3 2026", "status": "PLANNED", "deliverable": "State Police pilot"}
        ]
        render_milestone_roadmap(slide=slide, rect=rect, phases=phases, theme=self.theme, tokens=self.tokens)
        self.assertGreater(len(slide.shapes), 7)


if __name__ == "__main__":
    unittest.main()
