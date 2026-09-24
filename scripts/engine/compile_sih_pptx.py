#!/usr/bin/env python3
"""
Compile SIH Championship PowerPoint Decks and PDFs
Compiles high-resolution 4K rendered slide visuals into native 16:9 widescreen OpenXML PPTX decks
and lossless 4K multi-page PDFs using PyMuPDF FlateDecode compression.
"""

import os
import sys
import shutil
from pptx import Presentation
from pptx.util import Inches
import pymupdf
from PIL import Image

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
PRESENTATIONS_DIR = os.path.join(ROOT_DIR, "specs", "presentations")
RENDERED_DIR = os.path.join(PRESENTATIONS_DIR, "rendered")

# Explicit whitelist: CHAKRA and BHEDAK only. Never include approved_chakra or approved_bhedak.
DECKS = [
    {
        "name": "CHAKRA",
        "output_pptx": os.path.join(PRESENTATIONS_DIR, "CHAKRA_SIH2026.pptx"),
        "output_pdf": os.path.join(PRESENTATIONS_DIR, "CHAKRA_SIH2026.pdf"),
        "slides_dir": os.path.join(RENDERED_DIR, "chakra"),
        "title": "CHAKRA v3.0 - Cross-Chain Threat Attribution & Transaction Tracing System",
        "slides_count": 6
    },
    {
        "name": "BHEDAK",
        "output_pptx": os.path.join(PRESENTATIONS_DIR, "BHEDAK_SIH2026.pptx"),
        "output_pdf": os.path.join(PRESENTATIONS_DIR, "BHEDAK_SIH2026.pdf"),
        "slides_dir": os.path.join(RENDERED_DIR, "bhedak"),
        "title": "BHEDAK v3.0 - Automated Darknet Crawler & Sovereign De-Anonymization Platform",
        "slides_count": 6
    }
]

def compile_deck(deck_info):
    prs = Presentation()
    # 16:9 Widescreen (13.333" x 7.5")
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    slides_dir = deck_info["slides_dir"]
    output_pptx = deck_info["output_pptx"]
    output_pdf = deck_info["output_pdf"]

    print(f"\n[Compiler] Processing {deck_info['name']} ({deck_info['slides_count']} slides from {slides_dir})...")
    
    # 1. Compile PPTX with native lossless 4K image embedding
    for i in range(1, deck_info["slides_count"] + 1):
        slide_img = os.path.join(slides_dir, f"slide_{i}.png")
        if not os.path.exists(slide_img):
            raise FileNotFoundError(f"Missing rendered slide: {slide_img}")
        
        # Verify resolution is 4K
        with Image.open(slide_img) as img:
            w, h = img.size
            if (w, h) != (3840, 2160):
                print(f"  [Warning] Slide {i} dimension is {w}x{h}, expected (3840, 2160).")
            else:
                print(f"  Slide {i} verified 4K UHD ({w}x{h}).")

        slide = prs.slides.add_slide(blank_layout)
        # Add high-resolution image filling exact 16:9 slide boundaries
        slide.shapes.add_picture(slide_img, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)

    prs.save(output_pptx)
    pptx_size_mb = os.path.getsize(output_pptx) / (1024 * 1024)
    print(f"[Compiler] Saved 4K PPTX: {output_pptx} ({pptx_size_mb:.2f} MB)")

    # 2. Compile Lossless 4K Multi-Page PDF via PyMuPDF FlateDecode
    pdf_doc = pymupdf.open()
    for i in range(1, deck_info["slides_count"] + 1):
        slide_img = os.path.join(slides_dir, f"slide_{i}.png")
        imgdoc = pymupdf.open(slide_img)
        pdfbytes = imgdoc.convert_to_pdf()
        imgpdf = pymupdf.open("pdf", pdfbytes)
        pdf_doc.insert_pdf(imgpdf)
        imgdoc.close()
        imgpdf.close()

    pdf_doc.save(output_pdf, deflate=True, garbage=4)
    pdf_doc.close()
    pdf_size_mb = os.path.getsize(output_pdf) / (1024 * 1024)
    print(f"[Compiler] Saved 4K PDF:  {output_pdf} ({pdf_size_mb:.2f} MB)")

    # 3. Synchronize mirrored copies into rendered folder for ease of access
    deck_folder = slides_dir
    rendered_root = RENDERED_DIR

    # Mirrored in specs/presentations/rendered/<deck>/
    shutil.copy2(output_pptx, os.path.join(deck_folder, f"{deck_info['name']}_SIH2026.pptx"))
    shutil.copy2(output_pdf, os.path.join(deck_folder, f"{deck_info['name']}_SIH2026.pdf"))

    # Mirrored in specs/presentations/rendered/
    shutil.copy2(output_pptx, os.path.join(rendered_root, f"{deck_info['name']}_SIH2026.pptx"))
    shutil.copy2(output_pdf, os.path.join(rendered_root, f"{deck_info['name']}_SIH2026.pdf"))
    print(f"[Compiler] Mirrored {deck_info['name']} PPTX and PDF to rendered directories.")

def main():
    print("=" * 70)
    print("SIH 2026 Presentation Subsystem - 4K PPTX & Lossless PDF Compilation")
    print("=" * 70)

    for deck in DECKS:
        compile_deck(deck)

    print("\n[Compiler] All presentations successfully compiled in 4K resolution!")

if __name__ == "__main__":
    main()
